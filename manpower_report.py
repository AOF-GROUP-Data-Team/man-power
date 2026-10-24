# ============================================================
# Branch Manpower Report — Shawarma Classic (567430) + Garatis (671657)
# GitHub Actions version. Latest submission per branch -> cleaned manpower
# -> xlsx -> emailed as attachment.
# Runs on SUNDAY (KSA) so the week covered is the full Sunday -> Saturday week
# that just ended, including Saturday's submissions.
# Secrets: ZENPUT_API_TOKEN, GMAIL_APP_PASSWORD
# ============================================================
import os, sys, smtplib, requests, re, ast
from datetime import datetime, timedelta
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email.utils import formataddr
from email import encoders
import pytz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

API_TOKEN     = os.environ.get("ZENPUT_API_TOKEN", "")
GMAIL_USER    = "aof.group.auto@gmail.com"
GMAIL_PASS    = os.environ.get("GMAIL_APP_PASSWORD", "")
SENDER_NAME   = "Business Intelligence"
TO_EMAILS     = ["m.alhuaydar@aofgroup.com","a.omara@aofgroup.com","s.alharbi@aofgroup.com"]
CC_EMAILS     = ["a.alsalem@aofgroup.com","o.salahaddin@aofgroup.com"]
MODE          = os.environ.get("MODE", "auto").strip().lower()   # auto | force | test
RUN_WEEKDAY   = 6      # Sunday (Mon=0 ... Sun=6), KSA time
SEND_EMAIL    = MODE != "test"
BASE          = "https://www.zenput.com/api/v3/submissions/"
HEADERS       = {"X-API-TOKEN": API_TOKEN}
TZ            = pytz.timezone("Asia/Riyadh")
PAGE          = 100
MAX_PAGES     = 200     # safety cap: 200 x 100 = 20,000 submissions per template
LINK_BASE     = "https://ops-execution.crunchtime.com/reports/#form_id/"

TEMPLATES = {567430: "Shawarma Classic", 671657: "Garatis"}

MIN_ID_LEN          = 2    # employee IDs shorter than this (after removing leading zeros) are ignored
HISTORY_WEEKS       = 26   # weeks of transfer history kept in the Transfers sheet

# Expected branches: brand -> {code: label}
CLASSIC = {
    "B01": "NURUH", "B02": "KHRUH", "B04": "NSRUH", "B05": "RWRUH",
    "B06": "DARUH", "B07": "LBRUH", "B08": "SWRUH", "B09": "AZRUH", "B10": "SHRUH",
    "B11": "NRRUH", "B12": "TWRUH", "B13": "AQRUH", "B14": "RBRUH", "B15": "NDRUH",
    "B16": "BDRUH", "B17": "QRRUH", "B18": "TKRUH", "B19": "MURUH", "B21": "KRRUH",
    "B22": "OBJED", "B23": "SLAHS", "B24": "SFJED", "B25": "RWAHS", "B26": "HAJED",
    "B27": "SARUH", "B28": "MAJED", "B30": "QARUH", "B31": "ANRUH", "B32": "FYJED",
    "B33": "HIRJED", "B34": "URRUH", "B35": "IRRUH", "B36": "PSJED", "B37": "SHWMAK",
    "B38": "UHDMM", "B39": "HSRUH", "B40": "MZDMM",
}
GARATIS = {
    "QB01": "السويدي", "QB02": "النرجس", "QB03": "اليرموك", "QB04": "العقيق",
    "QB05": "الخافجي", "QB06": "الخليج", "QB07": "الندوة", "QB08": "المنصورة", "QB09": "النزهة",
}
GARATIS_LOC_IDS = {
    "2232755": "QB01", "2235805": "QB02", "2254072": "QB03", "2256386": "QB04",
    "2268360": "QB05", "2270650": "QB06", "2274188": "QB07", "2276794": "QB08", "2282774": "QB09",
}

PLACEHOLDER_NAMES = {"no", "n", "na", "none", "nil", "null", "x", "xx", "nothing", "empty",
                     "vacant", "no name", "no employee", "not available", "not availble",
                     "not avilable", "no one", "nobody"}

# ---------------- helpers ----------------
def parse_dt(s):
    if not s: return None
    try: return datetime.fromisoformat(s.replace("Z", "+00:00")).astimezone(TZ)
    except Exception: return None

def text_of(v):
    if v is None: return ""
    if isinstance(v, list): return ", ".join(map(str, v))
    if isinstance(v, str) and v.strip().startswith("["):
        try:
            lst = ast.literal_eval(v)
            if isinstance(lst, list): return ", ".join(map(str, lst))
        except Exception: pass
    return str(v).strip()

def name_compact(s):
    return re.sub(r"[^a-z\u0600-\u06ff]", "", (s or "").lower())

def is_valid_name(s):
    n = re.sub(r"\s+", " ", re.sub(r"[^a-z\u0600-\u06ff ]", "", (s or "").lower())).strip()
    if len(n.replace(" ", "")) < 2 or n in PLACEHOLDER_NAMES: return False
    if n.startswith("not avail") or n.startswith("no cashier"): return False
    return True

def id_key(s):
    return re.sub(r"\D", "", s or "").lstrip("0") or None

def is_yes(v):
    return str(v).strip().lower() in ("true", "yes", "1")

def resolve_branch(brand, loc):
    name = re.sub(r"\s+", " ", (loc.get("name") or loc.get("display_name") or "")).strip()
    if brand == "Garatis":
        code = GARATIS_LOC_IDS.get(str(loc.get("id") or ""))
        if not code:
            m = re.search(r"\bQB\s?(\d{2})\b", name, re.I) or re.search(r"Garatis+\s*B(\d{2})", name, re.I)
            code = f"QB{m.group(1)}" if m else None
        return code if code in GARATIS else None
    m = re.search(r"\bB(\d{2})\b", name)
    code = f"B{m.group(1)}" if m else None
    if code in CLASSIC and name.upper().startswith(CLASSIC[code]):
        return code
    return None

# ---------------- fetch ----------------
def fetch(template_id):
    out, start = [], 0
    for _ in range(MAX_PAGES):
        r = requests.get(BASE, headers=HEADERS, params={
            "form_template_id": template_id, "limit": PAGE, "start": start}, timeout=60)
        print(f"[{template_id}] GET start={start} -> HTTP {r.status_code}")
        if r.status_code != 200:
            print(r.text[:500]); break
        data = r.json().get("data", [])
        if not data: break
        out.extend(data)
        if len(data) < PAGE: break
        start += PAGE
    return out

# ---------------- manpower count ----------------
def parse_employees(s):
    """-> (unique valid employees [(id_key or None, name)], cashier_extra(0/1), slots)"""
    emps, cur = [], None
    cashier_flag, cashier_name = None, ""
    for a in s.get("answers", []):
        tl = (a.get("title") or "").strip().lower()
        v = a.get("value")
        if tl.startswith("write ") and tl.endswith("employee"):
            cur = {"name": text_of(v), "id": ""}
            emps.append(cur)
        elif tl == "id number" and cur is not None and not cur["id"]:
            cur["id"] = text_of(v)
        elif tl.startswith("this number of employees with cashier"):
            cashier_flag = is_yes(v)
        elif tl == "write name of cashier":
            cashier_name = text_of(v)
    seen_ids, seen_names, valid = set(), set(), []
    for e in emps:
        if not is_valid_name(e["name"]): continue
        ik, nk = id_key(e["id"]), name_compact(e["name"])
        if ik and len(ik) < MIN_ID_LEN: ik = None
        if (ik and ik in seen_ids) or (not ik and nk in seen_names): continue
        valid.append((ik, e["name"].strip()))
        if ik: seen_ids.add(ik)
        seen_names.add(nk)
    extra = 1 if (cashier_flag is False and is_valid_name(cashier_name)
                  and name_compact(cashier_name) not in seen_names) else 0
    return valid, extra, len(emps)

def count_manpower(s):
    valid, extra, slots = parse_employees(s)
    return len(valid) + extra, slots

# Weeks are anchored to midnight of the run day (KSA), so on a Sunday run
# week 0 = previous Sunday 00:00 -> Saturday 23:59, regardless of the run hour.
RUN_TIME = datetime.now(TZ).replace(hour=0, minute=0, second=0, microsecond=0)
def week_index(dt):
    """0 = the 7 days before the run day, 1 = the 7 days before that, ..."""
    return max(0, int((RUN_TIME - dt).total_seconds() // (7 * 86400)))

def week_range(w):
    end = RUN_TIME - timedelta(days=7 * w)
    return (end - timedelta(days=7)).date(), (end - timedelta(days=1)).date()

# ---------------- run ----------------
_today = datetime.now(TZ)
if MODE == "auto" and _today.weekday() != RUN_WEEKDAY:
    print(f"Not Sunday in KSA ({_today.strftime('%A %Y-%m-%d')}) - nothing to do")
    sys.exit(0)
print(f"MODE={MODE} | {_today.strftime('%A %Y-%m-%d %H:%M')} KSA")
if not API_TOKEN:
    sys.exit("ZENPUT_API_TOKEN is not set")
rows = []
snap = {}      # (branch label, week index) -> (latest datetime, set of IDs)
first_seen = {} # (employee id, branch label, week index) -> earliest datetime listed
for tid, brand in TEMPLATES.items():
    subs = fetch(tid)
    print(f"[{tid}] {brand}: {len(subs)} submissions")
    latest = {}
    for s in subs:
        sm = s.get("smetadata") or {}
        code = resolve_branch(brand, sm.get("location") or {})
        dt = parse_dt(sm.get("date_submitted"))
        if not code or not dt:
            continue
        if code not in latest or dt > latest[code][0]:
            latest[code] = (dt, s)
        exp_map = CLASSIC if brand == "Shawarma Classic" else GARATIS
        label = f"{code} {exp_map[code]}"
        w = week_index(dt)
        ids = {ik for ik, _ in parse_employees(s)[0] if ik}
        if (label, w) not in snap or dt > snap[(label, w)][0]:
            snap[(label, w)] = (dt, ids)
        for ik in ids:
            k = (ik, label, w)
            if k not in first_seen or dt < first_seen[k]: first_seen[k] = dt
    expected = CLASSIC if brand == "Shawarma Classic" else GARATIS
    structure_warned = False
    for code in sorted(expected, key=lambda c: int(re.sub(r"\D", "", c))):
        label = f"{code} {expected[code]}"
        if code not in latest:
            rows.append([brand, label, None, "No submission", None])
            continue
        dt, s = latest[code]
        total, slots = count_manpower(s)
        if slots == 0 and not structure_warned:
            structure_warned = True
            print(f"\n!! [{tid}] no employee fields recognised in {label} — answer titles:")
            for a in s.get("answers", []):
                print(f"   {a.get('title')} / {a.get('field_type')} / {text_of(a.get('value'))[:60]}")
        rows.append([brand, label, total, dt.strftime("%Y-%m-%d %H:%M"),
                     LINK_BASE + str(s.get("legacy_submission_id") or "")])

# ---------------- transfers (ID only, week vs previous week) ----------------
# employee's branches per week = branches whose LATEST submission that week lists the ID
emp_weeks = {}
for (label, w), (dt, ids) in snap.items():
    for ik in ids:
        emp_weeks.setdefault(ik, {}).setdefault(w, set()).add(label)

transfers = []
for ik, weeks in emp_weeks.items():
    ws_sorted = sorted(weeks, reverse=True)            # oldest week first
    for i, (prev_w, cur_w) in enumerate(zip(ws_sorted, ws_sorted[1:])):
        if cur_w > HISTORY_WEEKS: continue
        before = weeks[ws_sorted[i - 1]] if i > 0 else set()   # week before prev_w
        old = weeks[prev_w] - weeks[cur_w]               # branches that dropped the ID
        # branches listing the ID now; a branch that already had the ID before the overlap is
        # its long-standing branch (the dropped one was a one-off typo), so it is not a transfer
        targets = {t for t in weeks[cur_w] - old if not (t in weeks[prev_w] and t in before)}
        if not old or not targets: continue               # no move
        for to in sorted(targets):
            # transfer date = first time the ID appeared at the new branch (may be the week before,
            # when the employee was briefly listed at both branches)
            seen = [first_seen[(ik, to, w)] for w in (cur_w, prev_w) if (ik, to, w) in first_seen]
            transfers.append({
                "id": ik, "week": cur_w,
                "from": " / ".join(sorted(old)),
                "last_from": max(snap[(o, prev_w)][0] for o in old),
                "to": to, "first_to": min(seen),
            })
transfers.sort(key=lambda t: (t["week"], t["to"], t["id"]))
week_transfers = [t for t in transfers if t["week"] == 0]
wk0_start, wk0_end = week_range(0)
print(f"\nEmployee IDs tracked: {len(emp_weeks)} | transfers (last {HISTORY_WEEKS} weeks): {len(transfers)} "
      f"| this week ({wk0_start} - {wk0_end}): {len(week_transfers)}")
for t in week_transfers:
    print(f"  ID {t['id']:>6}  {t['from']:<16} -> {t['to']:<16} on {t['first_to']:%Y-%m-%d}")

# ---------------- xlsx ----------------
wb = Workbook()
ws = wb.active
ws.title = "Manpower"
ws.append(["Brand", "Branch", "Total Manpower", "Last Submitted", "Submission Link"])
thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
for r in rows:
    ws.append(r[:4] + ["Open" if r[4] else ""])
    row = ws[ws.max_row]
    if r[4]:
        row[4].hyperlink, row[4].font = r[4], Font(color="0563C1", underline="single")
    else:
        for c in row: c.fill = PatternFill("solid", fgColor="D9D9D9")
ws.append([])
for brand in TEMPLATES.values():
    ws.append([brand, "TOTAL", sum(r[2] or 0 for r in rows if r[0] == brand), "", ""])
    for c in ws[ws.max_row]: c.font = Font(bold=True)
ws.append(["ALL", "TOTAL", sum(r[2] or 0 for r in rows), "", ""])
for c in ws[ws.max_row]: c.font = Font(bold=True)

for c in ws[1]:
    c.fill, c.font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(horizontal="center", vertical="center")
        if c.value is not None: c.border = thin
for col, w in zip("ABCDE", [18, 20, 15, 18, 16]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"

def transfer_sheet(title, items):
    sh = wb.create_sheet(title)
    sh.append(["Week", "Employee ID", "From Branch", "Last Seen at Old Branch", "To Branch", "Transfer Date"])
    last_week = None
    for t in items:
        a, b = week_range(t["week"])
        sh.append([f"{a} - {b}", t["id"], t["from"], t["last_from"].strftime("%Y-%m-%d"),
                   t["to"], t["first_to"].strftime("%Y-%m-%d")])
        if t["week"] % 2:   # shade alternate weeks
            for c in sh[sh.max_row]: c.fill = PatternFill("solid", fgColor="EEF2F8")
    if not items:
        sh.append(["No transfers"])
    for c in sh[1]:
        c.fill, c.font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sh.iter_rows(min_row=2):
        for c in row: c.border, c.alignment = thin, Alignment(horizontal="center", vertical="center")
    for col, w in zip("ABCDEF", [24, 12, 22, 18, 22, 14]):
        sh.column_dimensions[col].width = w
    sh.freeze_panes = "A2"

transfer_sheet("Transfers This Week", week_transfers)
transfer_sheet("Transfers History", transfers)

now = datetime.now(TZ)
fname = f"manpower_report_{now.strftime('%Y-%m-%d')}.xlsx"
wb.save(fname)
print(f"\nSaved {fname}")
for brand in TEMPLATES.values():
    br = [r for r in rows if r[0] == brand]
    print(f"{brand}: {sum(1 for r in br if r[4])}/{len(br)} branches submitted, "
          f"total manpower {sum(r[2] or 0 for r in br)}")


# ---------------- email ----------------
LOWEST_N = 5          # branches shown in the "lowest manpower" table per brand
BRAND_AR = {"Shawarma Classic": "شاورما كلاسيك", "Garatis": "قراطيس"}

TD = "padding:6px 12px;border:1px solid #cccccc;text-align:center"
TH = "padding:8px 12px;background:#1F3864;color:#ffffff;text-align:center"

def lowest_table(brand):
    br = sorted([r for r in rows if r[0] == brand and r[4]], key=lambda r: (r[2], r[1]))[:LOWEST_N]
    if not br:
        return "<p>لا توجد بيانات.</p>"
    body = "".join(
        f"<tr><td style='{TD}'>{i}</td><td style='{TD}' dir='auto'>{r[1]}</td>"
        f"<td style='{TD};color:#C00000'><b>{r[2]}</b></td><td style='{TD}' dir='ltr'>{r[3]}</td></tr>"
        for i, r in enumerate(br, 1))
    return (f"<table style='border-collapse:collapse;margin-bottom:18px'>"
            f"<tr><th style='{TH}'>#</th><th style='{TH}'>الفرع</th><th style='{TH}'>عدد الموظفين</th>"
            f"<th style='{TH}'>آخر تسليم</th></tr>{body}</table>")

def transfers_html():
    title = f"الموظفون المنقولون هذا الأسبوع (<span dir='ltr'>{wk0_start} - {wk0_end}</span>)"
    if not week_transfers:
        return f"<h3 style='color:#1F3864;margin:18px 0 8px'>{title}</h3><p>لا توجد تنقلات هذا الأسبوع.</p>"
    body = "".join(
        f"<tr><td style='{TD}'>{i}</td><td style='{TD}'>{t['id']}</td>"
        f"<td style='{TD}' dir='auto'>{t['from']}</td><td style='{TD}' dir='auto'>{t['to']}</td>"
        f"<td style='{TD}' dir='ltr'>{t['first_to']:%Y-%m-%d}</td></tr>"
        for i, t in enumerate(week_transfers, 1))
    return (f"<h3 style='color:#1F3864;margin:18px 0 8px'>{title} - {len(week_transfers)} موظف</h3>"
            f"<table style='border-collapse:collapse'><tr><th style='{TH}'>#</th><th style='{TH}'>الرقم الوظيفي</th>"
            f"<th style='{TH}'>من فرع</th><th style='{TH}'>إلى فرع</th><th style='{TH}'>تاريخ النقل</th></tr>{body}</table>")

def build_html():
    summary = ""
    for brand in TEMPLATES.values():
        br = [r for r in rows if r[0] == brand]
        missing = [r[1] for r in br if not r[4]]
        summary += (f"<tr><td style='{TD}'>{BRAND_AR.get(brand, brand)}</td>"
                    f"<td style='{TD}'>{sum(1 for r in br if r[4])} / {len(br)}</td>"
                    f"<td style='{TD}'><b>{sum(r[2] or 0 for r in br)}</b></td>"
                    f"<td style='{TD}' dir='auto'>{'، '.join(missing) if missing else '-'}</td></tr>")
    summary += (f"<tr><td style='{TD}'><b>الإجمالي</b></td><td style='{TD}'></td>"
                f"<td style='{TD}'><b>{sum(r[2] or 0 for r in rows)}</b></td><td style='{TD}'></td></tr>")
    lowest = "".join(
        f"<h3 style='color:#1F3864;margin:18px 0 8px'>أقل {LOWEST_N} فروع في عدد الموظفين - {BRAND_AR.get(b, b)}</h3>{lowest_table(b)}"
        for b in TEMPLATES.values())
    return f"""<html><body dir="rtl" style="font-family:Tahoma,Arial,sans-serif;font-size:14px;color:#222;text-align:right">
<p>السلام عليكم ورحمة الله وبركاته،</p>
<p>مرفق لكم تقرير القوى العاملة في الفروع بتاريخ <span dir="ltr">{now.strftime('%Y-%m-%d')}</span>،
بناءً على آخر تسليم لكل فرع.</p>
<h3 style="color:#1F3864;margin:18px 0 8px">ملخص العلامات التجارية</h3>
<table style="border-collapse:collapse">
<tr><th style="{TH}">العلامة التجارية</th><th style="{TH}">الفروع المسلّمة</th>
<th style="{TH}">إجمالي الموظفين</th><th style="{TH}">فروع لم تسلّم</th></tr>
{summary}
</table>
{lowest}
{transfers_html()}
<p>تفاصيل جميع الفروع وروابط التسليم وسجل التنقلات في الملف المرفق.</p>
<p>مع خالص التحية،<br>إدارة ذكاء الأعمال</p>
</body></html>"""

def send_email(path):
    if not GMAIL_PASS:
        sys.exit("GMAIL_APP_PASSWORD is not set")
    msg = MIMEMultipart()
    msg["From"] = formataddr((SENDER_NAME, GMAIL_USER))
    msg["To"] = ", ".join(TO_EMAILS)
    msg["Cc"] = ", ".join(CC_EMAILS)
    msg["Subject"] = f"تقرير القوى العاملة في الفروع - {now.strftime('%Y-%m-%d')}"
    msg.attach(MIMEText(build_html(), "html", "utf-8"))
    with open(path, "rb") as f:
        part = MIMEBase("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f'attachment; filename="{os.path.basename(path)}"')
    msg.attach(part)
    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as s:
        s.login(GMAIL_USER, GMAIL_PASS)
        s.sendmail(GMAIL_USER, TO_EMAILS + CC_EMAILS, msg.as_string())
    print(f"Email sent to {TO_EMAILS} cc {CC_EMAILS}")

if SEND_EMAIL:
    send_email(fname)
else:
    print("MODE=test -> email skipped")
